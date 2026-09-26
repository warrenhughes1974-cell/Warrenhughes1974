"""Issue #118 — every quikridr.MUWCLASS must be a valid QLA UW class (no orphans)."""
from __future__ import annotations

import csv
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "QLA_Migration" / "Output"
RIDR = OUT / "quikridr.csv"
PLUW = OUT / "rates" / "QuikPlUw.csv"
UWPO = OUT / "rates" / "QuikUwpo.csv"
CROSSWALK = ROOT / "plan_analysis/source_data/crosswalk/Policy Form Crosswalk 5.22.26.xlsx"

APPROVED = {"00", "ST", "PR", "SM", "BL", "NT", "PQ"}
FORBIDDEN_LEGACY = {"NS", "T", "R", "Q", "N", "B", "S", "P", "M"}

UAT = {
    "9011189929C": "BL",
    "9011190516C": "SM",
    "9011193156C": "PR",
    "9011059291C": "ST",
    "9011052719C": "PR",
    "9011206462C": "NT",
    "9011208194C": "ST",
    "9011207210C": "PQ",
    "9011215903C": "PR",
    "9010360290C": "00",
}


def check_l10_family_coverage() -> bool:
    """L10_PLANS is hand-enumerated; prove it still covers every L10* coverage in the crosswalk.

    QLAdmin plan codes hide the family (L10 SPSWP -> 910SWP, L10 WP CDT -> 9CDTWP), so a
    missing entry silently maps LifePRO S to Standard instead of Smoker on that product.
    """
    if not CROSSWALK.is_file():
        print("WARN: Policy Form Crosswalk not found - skipped L10 family coverage check")
        return True
    try:
        import openpyxl

        sys.path.insert(0, str(ROOT))
        from qla_core import rate_dbf_schema as S
    except ImportError as exc:
        print(f"WARN: skipped L10 family coverage check ({exc})")
        return True

    wb = openpyxl.load_workbook(CROSSWALK, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    expected = set()
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        cov = row[0] if len(row) > 0 else None
        plan = row[2] if len(row) > 2 else None
        if cov and plan and str(cov).strip().upper().startswith("L10"):
            expected.add(str(plan).strip())
    missing = sorted(expected - set(S.L10_PLANS))
    if missing:
        print(f"  FAIL: L10_PLANS missing {len(missing)} L10 coverage plan(s): {missing}")
        return False
    print(f"  OK: L10_PLANS covers all {len(expected)} L10 crosswalk plans")
    return True


def main() -> int:
    ok = True
    if not RIDR.is_file():
        print(f"FAIL: missing {RIDR}")
        return 1

    if not check_l10_family_coverage():
        ok = False

    pluw_by_plan: dict[str, set[str]] = defaultdict(set)
    if PLUW.is_file():
        with PLUW.open(newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                pluw_by_plan[(r.get("PLAN") or "").strip()].add((r.get("UWCODE") or "").strip())

    uwpo_codes: set[str] = set()
    if UWPO.is_file():
        with UWPO.open(newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                uwpo_codes.add((r.get("UWCODE") or "").strip())
        if "NS" in uwpo_codes:
            print("FAIL: QuikUwpo still contains NS (should be dropped)")
            ok = False
        else:
            print("OK: QuikUwpo has no NS")
        extra = uwpo_codes - APPROVED
        if extra:
            print(f"FAIL: QuikUwpo unexpected codes: {sorted(extra)}")
            ok = False

    counts = Counter()
    orphans = []
    membership_miss = []
    uat_got = {}
    with RIDR.open(newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            pol = (r.get("MPOLICY") or "").strip()
            plan = (r.get("MPLAN") or "").strip()
            phase = (r.get("MPHASE") or "").strip()
            muw = (r.get("MUWCLASS") or "").strip()
            counts[muw or "(blank)"] += 1
            if not muw or muw not in APPROVED:
                orphans.append((pol, plan, phase, muw))
            elif pluw_by_plan and plan in pluw_by_plan and muw not in pluw_by_plan[plan]:
                membership_miss.append((pol, plan, muw, sorted(pluw_by_plan[plan])))
            if pol in UAT and phase in ("1", "01"):
                uat_got[pol] = muw

    print(f"Issue #118 MUWCLASS validation (rows={sum(counts.values())})")
    print("  MUWCLASS counts:", dict(counts.most_common()))
    if orphans:
        print(f"  FAIL: {len(orphans)} orphan/invalid MUWCLASS (sample):")
        for row in orphans[:15]:
            print(f"    {row}")
        ok = False
    else:
        print("  OK: every MUWCLASS in approved domain", sorted(APPROVED))

    # Membership miss is WARN for plans that only have default keys — still report
    if membership_miss:
        print(f"  WARN: {len(membership_miss)} rows MUWCLASS not in QuikPlUw for plan (sample):")
        for row in membership_miss[:10]:
            print(f"    {row}")

    for pol, expect in UAT.items():
        got = uat_got.get(pol)
        if got != expect:
            print(f"  FAIL: UAT {pol} MUWCLASS={got!r} expected {expect!r}")
            ok = False
        else:
            print(f"  OK: UAT {pol}={got}")

    print("RESULT:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
